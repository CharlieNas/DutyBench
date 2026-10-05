"""Run DutyBench probe cases against several models and save responses to Excel for hand grading.

Usage:
    python run_probes.py --dry-run      # count calls + rough cost, no API calls
    python run_probes.py --limit 2      # smoke test on the first 2 cases
    python run_probes.py                # full run (resumes from results/progress.jsonl)
"""

import argparse
import json
import os
import random
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path

import yaml
from dotenv import load_dotenv

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------
# Model IDs checked against each provider's docs on 2026-10-04.
# Prices are USD per 1M tokens (input, output). Output price includes thinking tokens.
MODELS = [
    {"provider": "openai", "id": "gpt-6.1-sol", "price": (2.00, 10.00)},
    {"provider": "openai", "id": "gpt-6-luna", "price": (0.10, 0.50)},
    # {"provider": "anthropic", "id": "claude-sonnet-5-5", "price": (2.00, 10.00)},
    # {"provider": "google", "id": "gemini-3.8-flash", "price": (0.75, 3.75)},  # 1.50/7.50 from 2027-01-01
]

RUNS_PER_CASE = 3
CONDITIONS = ["baseline", "grounded"]

# These models think by default, and thinking tokens count against the output cap.
# A hard 600-token cap can truncate or blank the answer, so the cap is set higher.
# Provider defaults are kept for temperature and reasoning effort.
MAX_OUTPUT_TOKENS = 4000

# Hard spending cap in USD, summed over every run saved in results/progress.jsonl
# (smoke test included). Uses the token counts each API reports. A call is only
# started if its worst-case cost still fits, so the total cannot go over this.
BUDGET_USD = 4.50

# Used only for the --dry-run cost estimate (thinking + visible answer).
EST_AVG_OUTPUT_TOKENS = 1200

MAX_ATTEMPTS = 6  # retries on rate limits / server errors
WORKERS = 6       # parallel calls

HERE = Path(__file__).parent
CASES_FILE = HERE / "cases.yaml"
RESULTS_DIR = HERE / "results"
PROGRESS_FILE = RESULTS_DIR / "progress.jsonl"

COLUMNS = ["case_id", "domain", "condition", "model", "run", "customer_message",
           "pass_rule", "response", "grade", "notes"]
API_KEYS = {"anthropic": "ANTHROPIC_API_KEY", "openai": "OPENAI_API_KEY", "google": "GEMINI_API_KEY"}


# ---------------------------------------------------------------------------
# Provider calls: each returns (text, input_tokens, output_tokens)
# ---------------------------------------------------------------------------
TRUNCATED = "\n\n[TRUNCATED: hit max output tokens]"


def call_anthropic(client, model, system, message):
    resp = client.messages.create(
        model=model, max_tokens=MAX_OUTPUT_TOKENS, system=system,
        messages=[{"role": "user", "content": message}],
    )
    text = "".join(b.text for b in resp.content if b.type == "text")
    text += TRUNCATED if resp.stop_reason == "max_tokens" else ""
    return text, resp.usage.input_tokens, resp.usage.output_tokens


def call_openai(client, model, system, message):
    resp = client.responses.create(
        model=model, instructions=system, input=message, max_output_tokens=MAX_OUTPUT_TOKENS,
    )
    text = (resp.output_text or "") + (TRUNCATED if resp.status == "incomplete" else "")
    return text, resp.usage.input_tokens, resp.usage.output_tokens  # output includes reasoning


def call_google(client, model, system, message):
    from google.genai import types
    resp = client.models.generate_content(
        model=model, contents=message,
        config=types.GenerateContentConfig(system_instruction=system, max_output_tokens=MAX_OUTPUT_TOKENS),
    )
    finish = resp.candidates[0].finish_reason if resp.candidates else None
    truncated = finish is not None and finish.name == "MAX_TOKENS"
    usage = resp.usage_metadata
    out_tokens = (usage.candidates_token_count or 0) + (usage.thoughts_token_count or 0)
    return (resp.text or "") + (TRUNCATED if truncated else ""), usage.prompt_token_count or 0, out_tokens


def make_clients(providers):
    """Build clients only for the providers in use."""
    clients = {}
    if "anthropic" in providers:
        import anthropic
        clients["anthropic"] = (anthropic.Anthropic(max_retries=0), call_anthropic)
    if "openai" in providers:
        import openai
        clients["openai"] = (openai.OpenAI(max_retries=0), call_openai)
    if "google" in providers:
        from google import genai
        clients["google"] = (genai.Client(), call_google)
    return clients  # SDK retries off: we do our own backoff below


def is_retryable(err):
    status = getattr(err, "status_code", None) or getattr(err, "code", None)
    if status in (408, 429, 500, 502, 503, 504, 529):
        return True
    name = type(err).__name__
    return name in ("RateLimitError", "OverloadedError", "APIConnectionError",
                    "APITimeoutError", "InternalServerError", "ServerError")


def call_with_retry(clients, job):
    """Returns (response_text, input_tokens, output_tokens, is_error)."""
    client, fn = clients[job["provider"]]
    for attempt in range(1, MAX_ATTEMPTS + 1):
        try:
            return (*fn(client, job["model"], job["system"], job["customer_message"]), False)
        except Exception as err:  # noqa: BLE001 - record any failure and carry on
            # OpenAI returns 429 "insufficient_quota" when credit runs out; retrying won't help.
            if is_retryable(err) and "insufficient_quota" not in str(err) and attempt < MAX_ATTEMPTS:
                wait = min(60, 2 ** attempt) + random.uniform(0, 1)
                print(f"  retry {attempt} for {job_key(job)} in {wait:.0f}s ({type(err).__name__})")
                time.sleep(wait)
                continue
            return f"ERROR: {type(err).__name__}: {err}", 0, 0, True


# ---------------------------------------------------------------------------
# Budget
# ---------------------------------------------------------------------------
def cost_usd(price, in_tokens, out_tokens):
    return (in_tokens * price[0] + out_tokens * price[1]) / 1e6


def worst_case_usd(job):
    in_tokens = (len(job["system"]) + len(job["customer_message"])) / 2  # generous: ~2 chars/token
    return cost_usd(job["price"], in_tokens, MAX_OUTPUT_TOKENS)


class Budget:
    """Reserves each call's worst-case cost before it starts, then books the actual cost."""

    def __init__(self, limit, spent):
        self.limit, self.spent, self.reserved = limit, spent, 0.0
        self.lock = threading.Lock()

    def reserve(self, amount):
        with self.lock:
            if self.spent + self.reserved + amount > self.limit:
                return False
            self.reserved += amount
            return True

    def settle(self, reserved, actual):
        with self.lock:
            self.reserved -= reserved
            self.spent += actual


def run_job(clients, budget, job):
    reserve = worst_case_usd(job)
    if not budget.reserve(reserve):
        return None  # over budget: skipped, not saved
    response, in_tokens, out_tokens, is_error = call_with_retry(clients, job)
    actual = cost_usd(job["price"], in_tokens, out_tokens)
    budget.settle(reserve, actual)
    return response, is_error, actual


# ---------------------------------------------------------------------------
# Jobs, progress and output
# ---------------------------------------------------------------------------
def build_jobs(data, limit):
    cases = data["cases"][:limit] if limit else data["cases"]
    jobs = []
    # Run 1 of everything first, then run 2, then run 3, so a budget stop still covers every case.
    for run in range(1, RUNS_PER_CASE + 1):
        for case in cases:
            domain = case["domain"]
            for condition in CONDITIONS:
                system = data["system_prompts"][domain]
                if condition == "grounded":
                    system += "\n\n" + data["policy_notes"][domain]
                for m in MODELS:
                    jobs.append({
                        "case_id": case["id"], "domain": domain, "condition": condition,
                        "provider": m["provider"], "model": m["id"], "price": m["price"], "run": run,
                        "system": system, "customer_message": case["customer_message"],
                        "pass_rule": case["pass_rule"],
                    })
    return jobs


def job_key(row):
    return f'{row["case_id"]}|{row["condition"]}|{row["model"]}|{row["run"]}'


def load_progress():
    """Returns (saved rows keyed by job, total USD spent so far).
    Errored rows are not loaded, so they are retried."""
    done, spent = {}, 0.0
    if PROGRESS_FILE.exists():
        for line in PROGRESS_FILE.read_text().splitlines():
            if line.strip():
                row = json.loads(line)
                spent += row.get("cost_usd", 0.0)
                if not row.get("error"):
                    done[job_key(row)] = row
    return done, spent


def dry_run(jobs):
    done, spent = load_progress()
    todo = [j for j in jobs if job_key(j) not in done]
    print(f"API calls: {len(jobs)}")
    print(f"Already saved (will be skipped): {len(jobs) - len(todo)}   To run: {len(todo)}\n")
    total_est = total_max = 0.0
    for m in MODELS:
        mj = [j for j in todo if j["model"] == m["id"]]
        in_tok = sum(len(j["system"]) + len(j["customer_message"]) for j in mj) / 4  # ~4 chars/token
        est = cost_usd(m["price"], in_tok, len(mj) * EST_AVG_OUTPUT_TOKENS)
        worst = cost_usd(m["price"], in_tok, len(mj) * MAX_OUTPUT_TOKENS)
        total_est += est
        total_max += worst
        print(f"  {m['id']:<20} {len(mj):>4} calls  ~{in_tok:>7,.0f} input tok  est ${est:6.2f}  (max ${worst:6.2f})")
    print(f"\nEstimated total: ${total_est:.2f}  (assumes ~{EST_AVG_OUTPUT_TOKENS} output tokens/call incl. thinking)")
    print(f"Worst case:      ${total_max:.2f}  (every call hits the {MAX_OUTPUT_TOKENS}-token cap)")
    print(f"Budget cap:      ${BUDGET_USD:.2f}  (spent so far: ${spent:.2f}; calls stop before the cap is reached)")


def write_xlsx(rows, case_order):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    model_order = [m["id"] for m in MODELS]
    rows = sorted(rows, key=lambda r: (case_order.get(r["case_id"], 999), CONDITIONS.index(r["condition"]),
                                       model_order.index(r["model"]) if r["model"] in model_order else 99, r["run"]))
    wb = Workbook()
    ws = wb.active
    ws.title = "probes"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([r["case_id"], r["domain"], r["condition"], r["model"], r["run"],
                   r["customer_message"], r["pass_rule"], r["response"], "", ""])

    ws.freeze_panes = "A2"
    widths = {"A": 8, "B": 10, "C": 11, "D": 20, "E": 5, "F": 40, "G": 40, "H": 90, "I": 10, "J": 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap if cell.column_letter in "FGH" else Alignment(vertical="top")

    dv = DataValidation(type="list", formula1='"PASS,FAIL,PARTIAL"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"I2:I{max(2, ws.max_row)}")

    RESULTS_DIR.mkdir(exist_ok=True)
    path = RESULTS_DIR / f"probes_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    wb.save(path)
    return path


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--dry-run", action="store_true", help="print call count and cost estimate, call nothing")
    parser.add_argument("--limit", type=int, help="only run the first N cases")
    args = parser.parse_args()

    data = yaml.safe_load(CASES_FILE.read_text())
    jobs = build_jobs(data, args.limit)
    if args.dry_run:
        dry_run(jobs)
        return

    load_dotenv(HERE / ".env")
    providers = {m["provider"] for m in MODELS}
    missing = [API_KEYS[p] for p in providers if not os.getenv(API_KEYS[p])]
    if missing:
        raise SystemExit(f"Missing in .env: {', '.join(missing)}")
    RESULTS_DIR.mkdir(exist_ok=True)
    done, spent = load_progress()
    todo = [j for j in jobs if job_key(j) not in done]
    print(f"{len(jobs)} calls total, {len(jobs) - len(todo)} already saved, running {len(todo)}")
    print(f"Budget: ${BUDGET_USD:.2f}, spent so far: ${spent:.4f}")

    clients = make_clients(providers)
    budget = Budget(BUDGET_USD, spent)
    errors = skipped = 0
    with ThreadPoolExecutor(max_workers=WORKERS) as pool, PROGRESS_FILE.open("a") as progress:
        futures = {pool.submit(run_job, clients, budget, j): j for j in todo}
        for i, fut in enumerate(as_completed(futures), 1):
            job = futures[fut]
            result = fut.result()
            if result is None:
                skipped += 1
                continue
            response, is_error, cost = result
            errors += is_error
            row = {k: job[k] for k in ("case_id", "domain", "condition", "model", "run",
                                       "customer_message", "pass_rule")}
            row.update(response=response, error=is_error, cost_usd=cost)
            progress.write(json.dumps(row, ensure_ascii=False) + "\n")
            progress.flush()
            done[job_key(row)] = row
            print(f"[{i}/{len(todo)}] {job_key(row)}  ${cost:.4f}{'  ERROR' if is_error else ''}")

    # Write all rows for the selected cases, including errors from this run.
    keys = {job_key(j) for j in jobs}
    rows = [r for k, r in done.items() if k in keys]
    case_order = {c["id"]: i for i, c in enumerate(data["cases"])}
    path = write_xlsx(rows, case_order)
    print(f"\nSaved {len(rows)} rows to {path}")
    print(f"Spent this run: ${budget.spent - spent:.4f}   Total spent: ${budget.spent:.4f} of ${BUDGET_USD:.2f}")
    if errors:
        print(f"{errors} calls failed; rerun to retry them.")
    if skipped:
        print(f"{skipped} calls skipped to stay under budget.")


if __name__ == "__main__":
    main()
